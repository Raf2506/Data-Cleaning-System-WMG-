import io
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from invoice_cleaner import MappingLibrary, clean_dataframe, parse_invoice_listing, reports
from invoice_cleaner.cleaner import detect_brand, normalise_product, pack_type
from invoice_cleaner.names import split_customer
from invoice_cleaner.workbook import build_report_workbook, uom_bucket, with_uom_columns
from invoice_cleaner.parser import suggest_name_groups


def test_papadam_reports_under_cik_suri():
    assert detect_brand("PAPADAM SURI BULAT 100G X 10's X 20") == "CIK SURI"
    assert detect_brand("CIK SURI ASAM JAWA 900G") == "CIK SURI"
    assert detect_brand("TREVOR'S CREAMY HONEY BLEND") == "TREVOR"
    assert detect_brand("RASTO NACHO CHEESE SAUCE") == "RASTO"


def _fixture() -> io.BytesIO:
    """A miniature of the real export: metadata, page noise, headers, continuations."""
    rows = [
        ["Invoice Listing", None, None, None, None, None, None, None],
        ["Date : From 1/1/2026 to 31/7/2026", None, None, None, None, None, None, None],
        ["Company : All", None, None, None, None, None, None, None],
        ["ACME FOODS SDN BHD (123456-A)", None, None, None, None, None, None, None],
        ["Page 1 of 353", None, None, None, None, None, None, None],
        ["Doc. No", "Doc. Date", "Code", "Name", "Amount (RM)", None, None, None],
        ["IV-13371", "05/01/2026", "300-M0181", "ECONSAVE - AMPANG BARU", 1500.0, None, None, None],
        ["Seq", "GL Code", "Description", "Project", "Quantity", "UOM", "Unit Price", "Amount (RM)"],
        [1, "500-000", "RASTO CARBONARA MUSHROOM PASTA SAUCE", "----", 10, "CTN", 90.0, 900.0],
        [None, None, "350G X 12", None, None, None, None, None],
        [2, "500-000", "RASTO TOMATO BASIL", "----", 10, "CTN", 60.0, 600.0],
        ["ACME FOODS SDN BHD (123456-A)", None, None, None, None, None, None, None],
        ["Page 2 of 353", None, None, None, None, None, None, None],
        ["Doc. No", "Doc. Date", "Code", "Name", "Amount (RM)", None, None, None],
        ["IV-13372", "07/02/2026", "300-10042", "10068 AMPANG BARU", 400.0, None, None, None],
        ["Seq", "GL Code", "Description", "Project", "Quantity", "UOM", "Unit Price", "Amount (RM)"],
        [1, "500-000", "RASTO TOMATO BASIL", "----", 5, "CTN", 80.0, 400.0],
    ]
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False, header=False)
    buffer.seek(0)
    return buffer


def test_pipeline():
    parsed = parse_invoice_listing(_fixture())
    assert parsed.invoice_count == 2
    assert parsed.line_item_count == 3
    assert parsed.reported_range == ("1/1/2026", "31/7/2026")
    assert parsed.continuation_rows == 1
    # Continuation row was stitched onto the previous description.
    assert parsed.rows[0]["Product"].endswith("350G X 12")

    suggestions = suggest_name_groups(parsed.raw_names)
    assert suggestions["ECONSAVE - AMPANG BARU"] == "ECONSAVE"

    m = MappingLibrary()
    m.merge_suggestions(suggestions)
    # ECONSAVE reaches its rows two ways: the name, and the branch-code account
    # 300-10042 whose name is numeric.
    m.set_store("ECONSAVE", "ECONSAVE")
    m.set_store("300-10042", "ECONSAVE")

    frame = clean_dataframe(parsed, m)
    assert set(frame["OutletGroup"]) == {"ECONSAVE"}
    assert frame["Amount"].sum() == 1900.0


def test_rows_without_a_store_are_cleaned_not_dropped():
    """An empty keyword list must still produce a usable clean table.

    This is what a first upload from a new company looks like. Before, every row
    was out of scope and the whole upload came back empty.
    """
    parsed = parse_invoice_listing(_fixture())
    frame = clean_dataframe(parsed, MappingLibrary())
    assert len(frame) == 3
    assert set(frame["Mapping Status"]) == {"auto"}
    # "ECONSAVE - AMPANG BARU" splits into its chain and its branch, and the
    # numbered name "10068 AMPANG BARU" is an ECONSAVE outlet too.
    assert set(frame["OutletGroup"]) == {"ECONSAVE"}
    assert set(frame["Outlet"]) == {"AMPANG BARU"}
    assert frame["Amount"].sum() == 1900.0


def test_numbered_outlet_names_are_econsave():
    """"10058 KLEBANG" names an ECONSAVE outlet; the place is the branch."""
    m = MappingLibrary()
    assert m.group_and_branch("10058 KLEBANG", "300-10046") == ("ECONSAVE", "KLEBANG", "auto")
    assert m.group_and_branch("10106 BATU GAJAH", "300-10106") == ("ECONSAVE", "BATU GAJAH", "auto")


def test_branch_keywords_do_not_reach_outside_their_chains():
    """A "CASH & CARRY" branch rule must not rename MUTAIYAS CASH & CARRY."""
    m = MappingLibrary()
    m.set_code("CASH & CARRY", "CASH & CARRY")     # a Borong Din branch rule
    m.set_store("BORONG DIN", "BORONG DIN")
    # Borong Din is in the keyword list, so its own branch rule applies.
    assert m.group_and_branch("BORONG DIN AS CASH & CARRY", "300-B0011")[:2] == (
        "BORONG DIN", "CASH & CARRY")
    # Mutaiyas is not, so it keeps its own name and its account code.
    assert m.group_and_branch("MUTAIYAS CASH & CARRY SDN.BHD.", "300-M0126") == (
        "MUTAIYAS CASH & CARRY", "300-M0126", "auto")


def test_exclude_keyword_drops_rows():
    """The escape hatch for staff and internal accounts."""
    m = MappingLibrary()
    m.set_store("ECONSAVE", MappingLibrary.EXCLUDE)
    frame = clean_dataframe(parse_invoice_listing(_fixture()), m)
    excluded = frame[frame["Mapping Status"] == "excluded"]
    assert len(excluded) == 2
    assert set(excluded["OutletGroup"]) == {MappingLibrary.OUT_OF_SCOPE}


def test_pack_type_and_product_normalisation():
    """Unit codes become words, and one SKU stays one SKU."""
    assert pack_type("CTN") == "Carton"
    assert pack_type("PCS") == "Pieces"
    assert pack_type("unit") == "Unit"
    assert pack_type("WIDGET") == "WIDGET"  # unknown codes pass through

    assert normalise_product("RASTO KOREAN SAUCE 250GX 24") == "RASTO KOREAN SAUCE 250G X 24"
    assert normalise_product("RASTO NACHOS CHEESE SAUCES 250G X 24") == "RASTO NACHO CHEESE SAUCE 250G X 24"
    assert normalise_product("CIK SURI FISH SAUCE 750ml X 12**FOC**") == "CIK SURI FISH SAUCE 750ML X 12"
    # "MAX 1KG" must not be split the way "250GX 24" is.
    assert normalise_product("RASTO MAYO MAX 1KG X 20") == "RASTO MAYO MAX 1KG X 20"


def _header_only_fixture() -> io.BytesIO:
    """An Invoice Listing exported with the line-item detail switched off."""
    rows = [
        ["Invoice Listing", None, None, None, None],
        ["Doc. No", "Doc. Date", "Code", "Name", "Amount (RM)"],
        ["IV-17050", "01/08/2026", "300-A0118", "AEON (KL RDC)", 164.46],
        ["IV-17053", "01/08/2026", "300-SNWG", "ST ROSYAM MART SDN BHD - SNWG", 5683.50],
    ]
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False, header=False)
    buffer.seek(0)
    return buffer


def test_export_without_line_detail_still_yields_invoice_totals():
    """Dropping it would lose a whole month of real sales."""
    parsed = parse_invoice_listing(_header_only_fixture())
    assert parsed.header_only is True
    assert parsed.invoice_count == 2
    assert sum(r["Amount"] for r in parsed.rows) == 5847.96
    # The product is unknown and says so, rather than being invented.
    assert {r["Product"] for r in parsed.rows} == {"(no line detail in export)"}


def test_unnamed_account_falls_back_to_its_invoice_code():
    """Some headers carry only a code; the code is the store, not a dropped row."""
    group, _, status = MappingLibrary().group_and_branch("", "300-BANGI")
    assert (group, status) == ("300-BANGI", "auto")


def _cleaned_export_fixture() -> io.BytesIO:
    """This tool's own XLSX export, re-uploaded. No document number anywhere."""
    rows = [
        ["CUSTOMER NAME", "OUTLET", "ITEM DESCRIPTIONS", "ITEM BRAND", "QUANTITY", "UNIT PRICE", "AMOUNT"],
        ["99 SPEED MART", "300-990006", "RASTO GARLIC SPREAD CHEESE 200G X 12", "RASTO", 95, 84, 7980],
        ["AOMORI MART", "JELAPANG", "SONGKHLA TOMYAM PASTE 227G X 12", "SONGKHLA", 10, 42, 420],
    ]
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False, header=False)
    buffer.seek(0)
    return buffer


def test_cleaned_export_can_be_uploaded_again():
    parsed = parse_invoice_listing(_cleaned_export_fixture())
    assert parsed.line_item_count == 2
    assert sum(r["Amount"] for r in parsed.rows) == 8400
    assert parsed.rows[0]["Raw Name"] == "99 SPEED MART"


def _three_row_frame():
    """A cleaned table with two stores, two months and three units of measure."""
    return pd.DataFrame(
        {
            "OutletGroup": ["ALPHA", "ALPHA", "BETA"],
            "Outlet": ["A1", "A1", "B1"],
            "Invoice No": ["IV-1", "IV-2", "IV-3"],
            "Date": pd.to_datetime(["2026-06-01", "2026-07-01", "2026-06-02"]),
            "Month": ["2026-06", "2026-07", "2026-06"],
            "Brand": ["RASTO", "RASTO", "MONT"],
            "Product": ["RASTO CHILI SAUCE 1KG X 20"] * 2 + ["MONT SYRUP 750ML X 12"],
            "Quantity": [10.0, 4.0, 6.0],
            "UOM": ["CTN", "PCS", "UNIT"],
            "Pack Type": ["Carton", "Pieces", "Unit"],
            "Unit Price": [80.0, 5.0, 20.0],
            "Amount": [800.0, 20.0, 120.0],
            "Code": ["300-A", "300-A", "300-B"],
            "Raw Name": ["ALPHA SDN BHD", "ALPHA SDN BHD", "BETA SDN BHD"],
            "Mapping Status": ["auto", "auto", "auto"],
        }
    )


def test_quantity_is_split_by_unit_of_measure():
    """Cartons, pieces and units must never be added into one number."""
    assert [uom_bucket(u) for u in ("CTN", "CTNe", "PCS", "UNIT", "BTL")] == [
        "CARTON", "CARTON", "PCS", "UNIT", "OTHER"]
    split = with_uom_columns(_three_row_frame())
    assert list(split["CARTON"]) == [10.0, 0.0, 0.0]
    assert list(split["PCS"]) == [0.0, 4.0, 0.0]
    assert list(split["UNIT"]) == [0.0, 0.0, 6.0]


def test_uom_mix_and_brand_ranking():
    frame = _three_row_frame()
    mix = {r["bucket"]: r for r in reports.uom_mix(frame)}
    assert mix["CARTON"]["quantity"] == 10.0
    assert mix["PCS"]["codes"] == ["PCS"]

    ranking = reports.brand_ranking(frame)
    assert [r["brand"] for r in ranking] == ["RASTO", "MONT"]        # best to worst
    assert round(ranking[0]["share"], 3) == round(820 / 940, 3)
    assert ranking[1]["unit"] == 6.0


def test_top_stores_per_month_ranks_within_each_month():
    months = reports.top_stores_per_month(_three_row_frame(), top_n=5)
    assert [m["month"] for m in months] == ["2026-06", "2026-07"]
    assert [s["store"] for s in months[0]["stores"]] == ["ALPHA", "BETA"]
    assert round(months[0]["stores"][0]["share"], 4) == round(800 / 920, 4)


def test_report_workbook_has_summary_index_and_a_sheet_per_store():
    book = load_workbook(io.BytesIO(build_report_workbook(_three_row_frame())))
    assert book.sheetnames[:2] == ["SUMMARY", "INDEX"]
    assert {"ALPHA", "BETA", "CLEAN DATA", "NOTES"} <= set(book.sheetnames)

    # CLEAN DATA carries the date, which the hand-built template lacked.
    headers = [c.value for c in book["CLEAN DATA"][1]]
    assert headers[:3] == ["MONTH", "DATE", "INVOICE NO"]
    assert book["CLEAN DATA"]["B2"].value is not None

    # ALPHA's two months sit side by side in their own column blocks.
    alpha = book["ALPHA"]
    assert alpha["A4"].value == "JUN 2026"
    assert alpha["M4"].value == "JUL 2026"  # 11 columns + a spacer


def test_customer_names_split_into_store_and_branch():
    """The conventions the user applies by hand in their own cleaned workbooks."""
    assert split_customer("99 SPEED MART SDN BHD") == ("99 SPEED MART", "")
    assert split_customer("AOMORI MART SDN BHD (JELAPANG)") == ("AOMORI MART", "JELAPANG")
    assert split_customer("SST APPAREL SDN BHD c/o Coffeehub") == ("SST APPAREL", "COFFEEHUB")
    assert split_customer("MUTAIYAS CASH & CARRY SDN.BHD.") == ("MUTAIYAS CASH & CARRY", "")
    # A bracketed registration number is not a branch; "(M)" is part of the name.
    assert split_customer("MOHAMED MEERA SAHIB (M) SDN.BHD. (405938-H)") == (
        "MOHAMED MEERA SAHIB (M)", "")


def _wide_fixture() -> io.BytesIO:
    """The real export's layout: 16 columns, fields spread by merged cells.

    Invoice fields land in 0/2/4/6/15 and line items in 0/1/3/8/10/11/12/13 —
    nothing like the narrow fixture above. Both must parse identically.
    """
    def row(pairs):
        cells = [None] * 16
        for index, value in pairs:
            cells[index] = value
        return cells

    rows = [
        row([(0, "Date"), (1, ":"), (2, "From 1/1/2026 to 31/7/2026")]),
        row([(0, "WEST MALAYAN GROUP SDN BHD"), (15, "Page 1 of 2")]),
        row([(0, "Doc. No"), (2, "Doc. Date"), (4, "Code"), (7, "Name"), (15, "Amount (RM)")]),
        row([(0, "IV-13371"), (2, datetime(2026, 1, 2)), (4, "300-M0181"),
             (6, "MYDIN MOHAMED HOLDINGS BHD"), (15, 2880.0)]),
        row([(0, "Seq"), (1, "GL Code"), (3, "Description"), (8, "Project"),
             (10, "Quantity"), (11, "UOM"), (12, "Unit Price"), (13, "Amount (RM)")]),
        row([(0, 1000), (1, "500-000"), (3, "RASTO GARLIC SPREAD BUTTER 200G X"), (8, "----"),
             (10, 10), (11, "CTN"), (12, 192.0), (13, 1920.0)]),
        row([(3, "24")]),
        row([(0, 2000), (1, "500-000"), (3, "RASTO MAYO GARLIC 250G X 24"), (8, "----"),
             (10, 5), (11, "CTN"), (12, 192.0), (13, 960.0)]),
        # Header that lost its "IV-" prefix because the cell was written as a number.
        row([(0, 16812), (2, datetime(2026, 7, 20)), (4, "300-N0115"),
             (6, "NSK GROCER (KL) SDN. BHD."), (15, 75.89)]),
        row([(0, "Seq"), (1, "GL Code"), (3, "Description"), (8, "Project"),
             (10, "Quantity"), (11, "UOM"), (12, "Unit Price"), (13, "Amount (RM)")]),
        row([(0, 1000), (1, "500-000"), (3, "CIK SURI ASAM JAWA XTRA 200G x 24"), (8, "----"),
             (10, 1), (11, "CTN"), (12, 79.68), (13, 75.89)]),
        # Item-summary block that closes the report: lone description cells here
        # must NOT be stitched onto the last invoice line.
        row([(1, "CSAJ200Gx24"), (3, "CIK SURI ASAM JAWA XTRA 200G x"), (7, 279), (8, "CTN"), (10, 19211.52)]),
        row([(3, "24")]),
        row([(14, "Total Item(s) :"), (15, 113)]),
    ]
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False, header=False)
    buffer.seek(0)
    return buffer


def test_wide_layout_is_parsed_by_content_not_position():
    parsed = parse_invoice_listing(_wide_fixture())
    assert parsed.invoice_count == 2
    assert parsed.line_item_count == 3
    assert parsed.reported_range == ("1/1/2026", "31/7/2026")

    first = parsed.rows[0]
    assert first["Invoice No"] == "IV-13371"
    assert first["Raw Name"] == "MYDIN MOHAMED HOLDINGS BHD"
    assert first["Code"] == "300-M0181"
    assert first["Quantity"] == 10
    assert first["UOM"] == "CTN"
    assert first["Unit Price"] == 192.0
    assert first["Amount"] == 1920.0
    # Continuation stitched despite the description sitting in column D.
    assert first["Product"] == "RASTO GARLIC SPREAD BUTTER 200G X 24"


def test_header_without_iv_prefix_is_not_read_as_a_line_item():
    parsed = parse_invoice_listing(_wide_fixture())
    assert "16812" in {r["Invoice No"] for r in parsed.rows}
    # Its line belongs to it, not to the invoice above.
    orphaned = [r for r in parsed.rows if r["Invoice No"] == "IV-13371"]
    assert sum(r["Amount"] for r in orphaned) == 2880.0


def test_trailing_summary_block_is_not_stitched():
    parsed = parse_invoice_listing(_wide_fixture())
    last = parsed.rows[-1]
    assert last["Product"] == "CIK SURI ASAM JAWA XTRA 200G x 24"
    assert parsed.continuation_rows == 1  # only the real one, not the summary's


def test_undetected_outlet_falls_back_to_the_invoice_code():
    """No name rule, no keyword rule — label by the code, and keep it flagged."""
    library = MappingLibrary()
    outlet, status = library.resolve("10094 KUBANG KERIAN", "300-10075")
    assert outlet == "300-10075"
    assert status == "unmapped"

    # With no code at all there is nothing better than the raw name.
    assert library.resolve("SOME GROCER", "") == ("SOME GROCER", "unmapped")


def test_keyword_matches_name_or_code_longest_first():
    library = MappingLibrary()
    library.set_code("KLUANG", "KLUANG")
    library.set_code("KLUANG PERDANA", "KLUANG PERDANA")
    library.set_code("SNWG", "SENAWANG")

    assert library.resolve("10043 KLUANG PERDANA", "300-1")[0] == "KLUANG PERDANA"
    assert library.resolve("10043 KLUANG", "300-1")[0] == "KLUANG"
    # Matched through the code rather than the name.
    assert library.resolve("10101 SD/CR", "300-SNWG01")[0] == "SENAWANG"


def test_store_matches_by_code_prefix_for_unnamed_accounts():
    """IKA-style: the export never names the chain, so a code fragment does it."""
    m = MappingLibrary(chain_keywords={"300-10": "ECONSAVE"})
    assert m.store_of("10068 AMPANG BARU", "300-10042") == "ECONSAVE"
    group, _, _ = m.group_and_branch("10068 AMPANG BARU", "300-10042")
    assert group == "ECONSAVE"


def test_clean_frame_carries_outlet_group():
    parsed = parse_invoice_listing(_fixture())
    m = MappingLibrary(chain_keywords={"ECONSAVE": "ECONSAVE", "300-10042": "ECONSAVE"})
    frame = clean_dataframe(parsed, m)
    assert "OutletGroup" in frame.columns
    assert set(frame["OutletGroup"]) == {"ECONSAVE"}


def test_chain_name_in_raw_gives_group_with_code_as_branch():
    """The SOON CHEONG case: chain named, branch not — group it, branch by code."""
    m = MappingLibrary(chain_keywords={"SOON CHEONG": "SOON CHEONG"})
    group, branch, status = m.group_and_branch("SOON CHEONG MARINE PRODUCT SDN BHD KL", "300-S0256")
    assert group == "SOON CHEONG"
    assert branch == "300-S0256"  # no outlet named, so the code stands in
    assert status == "mapped"


def test_store_in_name_wins_the_group_branch_only_labels():
    """ST ROSYAM MART (SEMENYIH) is SRI TERNAK's Semenyih branch."""
    m = MappingLibrary(chain_keywords={"ST": "SRI TERNAK"})
    m.set_code("SEMENYIH", "SEMENYIH")  # a Branch Outlet rule, not a store
    group, branch, _ = m.group_and_branch("ST ROSYAM MART (SEMENYIH)", "300-S0215")
    assert group == "SRI TERNAK"
    assert branch == "SEMENYIH"


def test_unknown_customer_is_auto_named_then_overridden_by_a_keyword():
    m = MappingLibrary(chain_keywords={"ST": "SRI TERNAK"})
    group, branch, status = m.group_and_branch("AEON CO. (M) BHD (KL RDC)", "300-A0118")
    assert (group, branch, status) == ("AEON CO. (M)", "KL RDC", "auto")
    # A keyword takes over the grouping and marks the row properly mapped.
    m.set_store("AEON", "AEON")
    assert m.group_and_branch("AEON CO. (M) BHD (KL RDC)", "300-A0118") == (
        "AEON", "KL RDC", "mapped")


def test_exclude_keeps_an_account_out_of_the_figures():
    m = MappingLibrary(chain_keywords={"STAFF": MappingLibrary.EXCLUDE})
    group, _, status = m.group_and_branch("STAFF SALES", "300-X0001")
    assert group == MappingLibrary.OUT_OF_SCOPE
    assert status == "excluded"


def test_chain_keyword_tolerates_plural_but_not_a_longer_word():
    m = MappingLibrary(chain_keywords={"CS BROTHER": "CS BROTHER", "ST": "SRI TERNAK"})
    assert m.chain_in_name("CS BROTHERS SDN BHD") == "CS BROTHER"
    assert m.chain_in_name("STAR GROCER SDN BHD") == ""  # ST must not leak into STAR


def test_store_with_a_named_branch():
    m = MappingLibrary(chain_keywords={"SOON CHEONG": "SOON CHEONG"})
    m.set_code("SG BULOH", "SG BULOH")
    group, branch, _ = m.group_and_branch("SOON CHEONG SG BULOH", "300-S0257")
    assert (group, branch) == ("SOON CHEONG", "SG BULOH")


def test_plain_names_map_to_themselves():
    """A name with no branch suffix is canonical already — not 'unmapped'."""
    suggestions = suggest_name_groups(["MYDIN MOHAMED HOLDINGS BHD", "ECONSAVE - AMPANG BARU"])
    assert suggestions["MYDIN MOHAMED HOLDINGS BHD"] == "MYDIN MOHAMED HOLDINGS BHD"
    assert suggestions["ECONSAVE - AMPANG BARU"] == "ECONSAVE"


def _tidy_fixture() -> io.BytesIO:
    """A clean one-row-per-line-item export under a real header row."""
    rows = [
        ["DocNo", "DocDate", "Code", "Name", "InvoiceAmount", "Seq", "GLCode",
         "Description", "Project", "Quantity", "UOM", "UnitPrice", "LineAmount"],
        ["IV-16481", "2026-07-01", "300-10025", "10026 BUTTERWORTH", 396.26, 1000,
         "500-000", "RASTO NACHO CHEESE SAUCE 1KG X 20", "----", 1, "CTN", 179.82, 171.26],
        ["IV-16481", "2026-07-01", "300-10025", "10026 BUTTERWORTH", 396.26, 2000,
         "500-000", "RASTO CHILI SAUCE 1KG X 20", "----", 3, "CTN", 75, 225],
        ["IV-16494", "2026-07-03", "300-S0205", "ST ROSYAM MART (SHAH ALAM)", 102.86, 1000,
         "500-000", "CIK SURI ASAM JAWA XTRA 200G x 36", "----", 1, "CTN", 108, 102.86],
    ]
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False, header=False)
    buffer.seek(0)
    return buffer


def test_tidy_table_is_parsed_by_column_name():
    parsed = parse_invoice_listing(_tidy_fixture())
    assert parsed.invoice_count == 2
    assert parsed.line_item_count == 3
    assert parsed.discarded_rows == 0
    first = parsed.rows[0]
    assert first["Invoice No"] == "IV-16481"
    assert first["Code"] == "300-10025"
    assert first["Product"] == "RASTO NACHO CHEESE SAUCE 1KG X 20"
    assert first["Amount"] == 171.26  # LineAmount, not InvoiceAmount
    assert first["Quantity"] == 1


def test_month_is_read_from_the_dates_not_a_metadata_block():
    """The tidy export has no 'Date : From ... to ...' header — dates carry it."""
    parsed = parse_invoice_listing(_tidy_fixture())
    assert {r["Month"] for r in parsed.rows} == {"2026-07"}
    assert parsed.date_from.date().isoformat() == "2026-07-01"
    assert parsed.date_to.date().isoformat() == "2026-07-03"


def _pascal_case_book_views(source: io.BytesIO) -> io.BytesIO:
    """Rewrite workbookView attributes the way non-Microsoft exporters do."""
    out = io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = data.replace(
                    b"<workbookView ",
                    b'<workbookView WindowWidth="28800" WindowHeight="12435" ',
                )
            zout.writestr(item, data)
    out.seek(0)
    return out


def test_pascalcase_workbook_attributes_are_repaired():
    """AutoCount-style exports use WindowWidth, not windowWidth; openpyxl refuses them."""
    broken = _pascal_case_book_views(_fixture())

    # Confirm the fixture really does break the underlying reader.
    broken.seek(0)
    try:
        pd.read_excel(broken, header=None, dtype=object)
        raise AssertionError("expected openpyxl to reject PascalCase attributes")
    except TypeError:
        pass

    broken.seek(0)
    parsed = parse_invoice_listing(broken)
    assert parsed.invoice_count == 2
    assert parsed.line_item_count == 3
    assert parsed.reported_range == ("1/1/2026", "31/7/2026")


def _renamed_tidy_fixture() -> io.BytesIO:
    """A re-exported (already cleaned) listing: InvoiceDate, not DocDate."""
    rows = [
        ["Item Description", "Seq", "GLCode", "Quantity", "UOM", "UnitPrice",
         "Amount", "InvoiceNo", "InvoiceDate", "CustomerCode", "CustomerName"],
        ["RASTO ROASTED SESAME SAUCE 250G X 24", 1000, "500-000", 10, "CTN",
         157.2, 1497.14, "IV12946", datetime(2026, 1, 1), "300-BANGI", "CS BROTHERS SDN BHD"],
        ["RASTO NACHO CHEESE SAUCE 250G X 24", 2000, "500-000", 2, "CTN",
         100.0, 200.0, "IV12947", datetime(2026, 6, 30), "300-BANGI", "CS BROTHERS SDN BHD"],
    ]
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False, header=False)
    buffer.seek(0)
    return buffer


def test_renamed_date_column_still_yields_months():
    """InvoiceDate/Invoice_Date must map, or every row loses its month."""
    parsed = parse_invoice_listing(_renamed_tidy_fixture())
    assert parsed.line_item_count == 2
    assert {r["Month"] for r in parsed.rows} == {"2026-01", "2026-06"}
    assert parsed.rows[0]["Code"] == "300-BANGI"
    assert parsed.rows[0]["Amount"] == 1497.14


if __name__ == "__main__":
    test_pipeline()
    test_papadam_reports_under_cik_suri()
    test_rows_without_a_store_are_cleaned_not_dropped()
    test_numbered_outlet_names_are_econsave()
    test_branch_keywords_do_not_reach_outside_their_chains()
    test_exclude_keyword_drops_rows()
    test_pack_type_and_product_normalisation()
    test_export_without_line_detail_still_yields_invoice_totals()
    test_unnamed_account_falls_back_to_its_invoice_code()
    test_cleaned_export_can_be_uploaded_again()
    test_quantity_is_split_by_unit_of_measure()
    test_uom_mix_and_brand_ranking()
    test_top_stores_per_month_ranks_within_each_month()
    test_report_workbook_has_summary_index_and_a_sheet_per_store()
    test_customer_names_split_into_store_and_branch()
    test_tidy_table_is_parsed_by_column_name()
    test_renamed_date_column_still_yields_months()
    test_month_is_read_from_the_dates_not_a_metadata_block()
    test_wide_layout_is_parsed_by_content_not_position()
    test_header_without_iv_prefix_is_not_read_as_a_line_item()
    test_trailing_summary_block_is_not_stitched()
    test_undetected_outlet_falls_back_to_the_invoice_code()
    test_keyword_matches_name_or_code_longest_first()
    test_store_matches_by_code_prefix_for_unnamed_accounts()
    test_clean_frame_carries_outlet_group()
    test_chain_name_in_raw_gives_group_with_code_as_branch()
    test_store_in_name_wins_the_group_branch_only_labels()
    test_unknown_customer_is_auto_named_then_overridden_by_a_keyword()
    test_exclude_keeps_an_account_out_of_the_figures()
    test_chain_keyword_tolerates_plural_but_not_a_longer_word()
    test_store_with_a_named_branch()
    test_plain_names_map_to_themselves()
    test_pascalcase_workbook_attributes_are_repaired()
    print("ok")
